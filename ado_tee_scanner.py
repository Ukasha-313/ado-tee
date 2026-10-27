"""
ADO-TEE Scanner Module
Handles file discovery, text extraction (.docx, .xlsx, .pdf, .txt),
SHA-256 duplicate calculation, and folder analysis.
"""

import hashlib
from pathlib import Path
import docx
import openpyxl
import pypdf

from ai_engine import AIEngine
from system_scanner import discover_files

# Master index output for manual scans
OUTPUT_FILE = Path("output/file_index.json")

# Supported document extensions
TARGET_EXTENSIONS = {".docx", ".xlsx", ".pdf", ".txt", ".md", ".csv"}

def calculate_hash(file_path: Path) -> str:
    """Calculate SHA-256 hash of a file for duplicate detection."""
    sha256_hash = hashlib.sha256()
    try:
        with open(file_path, "rb") as f:
            for byte_block in iter(lambda: f.read(4096), b""):
                sha256_hash.update(byte_block)
        return sha256_hash.hexdigest()
    except Exception:
        return ""

def extract_text(file_path: Path) -> str:
    """Extract readable text from PDF, DOCX, XLSX, TXT, or CSV files."""
    ext = file_path.suffix.lower()
    text = ""
    try:
        if ext in {".txt", ".md", ".csv"}:
            text = file_path.read_text(encoding="utf-8", errors="ignore")
        elif ext == ".docx":
            doc = docx.Document(file_path)
            text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
        elif ext == ".xlsx":
            wb = openpyxl.load_workbook(file_path, data_only=True)
            for sheet in wb.sheetnames:
                ws = wb[sheet]
                for row in ws.iter_rows(values_only=True):
                    row_text = " ".join([str(cell) for cell in row if cell is not None])
                    if row_text.strip():
                        text += row_text + "\n"
        elif ext == ".pdf":
            reader = pypdf.PdfReader(file_path)
            for page in reader.pages:
                extracted = page.extract_text()
                if extracted:
                    text += extracted + "\n"
    except Exception as error:
        print(f"[ERROR] Failed to extract text from {file_path.name}: {error}")
    
    return text

def scan_folder(folder_path: Path):
    """
    Scans a folder recursively, extracts text, checks for duplicates via hash,
    and analyzes content for categories and deadlines using the AI engine.
    """
    ai = AIEngine()
    files = []
    deadlines = []
    seen_hashes = set()

    discovered_files = discover_files(folder_path)

    for file_path in discovered_files:
        if file_path.suffix.lower() not in TARGET_EXTENSIONS:
            continue

        text = extract_text(file_path)
        file_hash = calculate_hash(file_path)

        # Duplicate detection check
        is_duplicate = False
        if file_hash and file_hash in seen_hashes:
            is_duplicate = True
        elif file_hash:
            seen_hashes.add(file_hash)

        # AI analysis
        analysis = ai.analyze_document(file_path.name, text)

        file_record = {
            "file_name": file_path.name,
            "full_path": str(file_path),
            "sha256": file_hash,
            "duplicate": is_duplicate,
            "text_length": len(text),
            **analysis
        }
        files.append(file_record)

        # Aggregate deadlines
        for deadline in analysis.get("deadlines", []):
            deadlines.append({
                "file": file_path.name,
                **deadline
            })

    return files, deadlines